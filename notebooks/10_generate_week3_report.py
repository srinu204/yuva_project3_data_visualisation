# ============================================================
# YUVA INTERNSHIP - WEEK 3
# Netflix Data Visualization and Reporting
# Generate the Final Word Report
# ============================================================

import pandas as pd

from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

# ------------------------------------------------------------
# 1. Project paths
# ------------------------------------------------------------

print(f"Running report script: {Path(__file__).resolve()}")

PROJECT_DIR = Path(__file__).resolve().parent.parent

DATA_PATH = PROJECT_DIR / "data" / "netflix_titles.csv"
CHART_DIR = PROJECT_DIR / "outputs" / "charts"
REPORT_DIR = PROJECT_DIR / "outputs" / "reports"

REPORT_DIR.mkdir(parents=True, exist_ok=True)

REPORT_PATH = REPORT_DIR / "Yuva_Week3_Data_Visualization_Report.docx"

# ------------------------------------------------------------
# 2. Load dataset and calculate summary statistics
# ------------------------------------------------------------

df = pd.read_csv(DATA_PATH)

total_titles = len(df)
total_columns = len(df.columns)

content_counts = df["type"].value_counts()

movie_count = int(content_counts.get("Movie", 0))
tv_count = int(content_counts.get("TV Show", 0))

movie_df = df[df["type"] == "Movie"].copy()

movie_df["duration_minutes"] = pd.to_numeric(
    movie_df["duration"].str.extract(r"(\d+)")[0],
    errors="coerce"
)

valid_durations = movie_df["duration_minutes"].dropna()

mean_duration = valid_durations.mean()
median_duration = valid_durations.median()

missing_values = df.isnull().sum()
total_missing = int(missing_values.sum())

# Country counts
country_df = df.dropna(subset=["country"]).copy()

country_df["country"] = country_df["country"].str.split(",")

country_df = country_df.explode("country")

country_df["country"] = country_df["country"].str.strip()

country_counts = country_df["country"].value_counts()

# Genre counts
genre_df = df[["listed_in"]].copy()

genre_df["listed_in"] = genre_df["listed_in"].str.split(",")

genre_df = genre_df.explode("listed_in")

genre_df["listed_in"] = genre_df["listed_in"].str.strip()

genre_counts = genre_df["listed_in"].value_counts()

# Content addition trends
df["date_added"] = pd.to_datetime(
    df["date_added"],
    format="mixed",
    errors="coerce"
)

valid_dates = df.dropna(subset=["date_added"]).copy()

yearly_additions = valid_dates.groupby(
    valid_dates["date_added"].dt.year
).size()

peak_year = int(yearly_additions.idxmax())
peak_additions = int(yearly_additions.max())

# Corrected movie-only correlation
# --------------------------------------------------
# Prepare movie data for correlation analysis
# --------------------------------------------------

# Select movies and create an independent copy
movie_df = df[df["type"] == "Movie"].copy()

# Convert date_added to datetime
movie_df["date_added"] = pd.to_datetime(
    movie_df["date_added"],
    format="mixed",
    errors="coerce"
)

# Extract the year in which the title was added to Netflix
movie_df["year_added"] = movie_df["date_added"].dt.year

# Extract numeric movie duration from values such as "90 min"
movie_df["duration_minutes"] = pd.to_numeric(
    movie_df["duration"]
    .astype("string")
    .str.extract(r"(\d+)", expand=False),
    errors="coerce"
)

# Ensure release_year is numeric
movie_df["release_year"] = pd.to_numeric(
    movie_df["release_year"],
    errors="coerce"
)

# Rebuild duration_minutes immediately before correlation analysis.
# This defensive step guarantees the column exists even if earlier code changes.
movie_df["duration_minutes"] = pd.to_numeric(
    movie_df["duration"]
    .astype("string")
    .str.extract(r"(\d+)", expand=False),
    errors="coerce"
)

# Confirm the required columns exist before selecting them.
required_correlation_columns = [
    "release_year", "year_added", "duration_minutes"
]
missing_correlation_columns = [
    column for column in required_correlation_columns
    if column not in movie_df.columns
]
if missing_correlation_columns:
    raise ValueError(
        "Could not prepare correlation columns: "
        + ", ".join(missing_correlation_columns)
        + f". Available columns: {list(movie_df.columns)}"
    )

# Select the required columns and remove missing values
correlation_data = movie_df[required_correlation_columns].dropna()

# Calculate the correlation matrix
correlation_matrix = correlation_data.corr()

# Display results in the terminal
print("\nMovie Correlation Analysis")
print("-" * 40)
print(f"Movies included in analysis: {len(correlation_data)}")
print("\nCorrelation Matrix:")
print(correlation_matrix.round(3))

# ------------------------------------------------------------
# 3. Create Word document
# ------------------------------------------------------------

doc = Document()

section = doc.sections[0]

section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

# Set default font
styles = doc.styles

styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(10.5)

styles["Title"].font.name = "Calibri"
styles["Title"].font.size = Pt(30)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(180, 0, 25)

for style_name in ["Heading 1", "Heading 2"]:
    styles[style_name].font.name = "Calibri"
    styles[style_name].font.color.rgb = RGBColor(180, 0, 25)

# ------------------------------------------------------------
# Helper functions
# ------------------------------------------------------------

def add_paragraph(text):
    doc.add_paragraph(text)


def add_bullet(text):
    doc.add_paragraph(text, style="List Bullet")


def add_heading(text, level=1):
    doc.add_heading(text, level=level)


def add_chart(filename, title, explanation, finding):
    chart_path = CHART_DIR / filename

    add_heading(title, level=2)

    if chart_path.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        paragraph.add_run().add_picture(
            str(chart_path),
            width=Inches(6.2)
        )
    else:
        add_paragraph(
            f"Chart image not found: {filename}. "
            "Generate the chart and run this report script again."
        )

    add_paragraph("Purpose and interpretation:")
    add_paragraph(explanation)

    add_paragraph("Key finding:")
    add_paragraph(finding)


def add_code_snippet(code):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Pt(8)


# ------------------------------------------------------------
# 4. Title page
# ------------------------------------------------------------

doc.add_paragraph("\n\n")

title = doc.add_heading(
    "DATA VISUALIZATION AND REPORTING",
    level=0
)

title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph(
    "Netflix Movies and TV Shows Analysis"
)

subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle.runs[0].font.size = Pt(18)
subtitle.runs[0].bold = True

doc.add_paragraph("\n")

p = doc.add_paragraph(
    "YUVA INTERNSHIP — WEEK 3 PROJECT"
)

p.alignment = WD_ALIGN_PARAGRAPH.CENTER

p.runs[0].bold = True

doc.add_paragraph("\n\n")

p = doc.add_paragraph(
    "Prepared using Python, Pandas, Matplotlib and Seaborn"
)

p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break()

# ------------------------------------------------------------
# 5. Abstract
# ------------------------------------------------------------

add_heading("Abstract")

add_paragraph(
    f"This project presents a data visualization study of the "
    f"Netflix Movies and TV Shows dataset containing {total_titles:,} "
    f"titles and {total_columns} columns. The objective is to transform "
    "raw catalog information into clear visualizations that communicate "
    "patterns, distributions, and relationships in the available data. "
    "Python libraries including Pandas, Matplotlib, and Seaborn were "
    "used for data preparation, statistical analysis, and visualization. "
    "The analysis includes country-wise content distribution, annual "
    "and monthly content additions, movie durations, genre frequencies, "
    "content-type proportions, and numerical correlations. The findings "
    "show that movies represent the majority of the recorded titles, "
    "International Movies is the most frequent category in the genre "
    "analysis, and movie durations are concentrated around the "
    "91–120-minute range. The report also discusses data limitations "
    "and the importance of interpreting correlations cautiously."
)

# ------------------------------------------------------------
# 6. Introduction and objectives
# ------------------------------------------------------------

add_heading("1. Introduction")

add_paragraph(
    "Data visualization is an important part of data analysis because "
    "it converts numerical and categorical information into visual "
    "representations that are easier to understand. Well-designed "
    "charts help analysts identify patterns, compare categories, "
    "observe changes over time, and communicate evidence-based findings."
)

add_paragraph(
    "This project uses Netflix catalog data to demonstrate a structured "
    "visualization workflow, from initial data exploration to chart "
    "creation, interpretation, and reporting."
)

add_heading("2. Project Objectives")

for item in [
    "Explore the structure and quality of the Netflix dataset.",
    "Identify the most frequently represented countries and genres.",
    "Analyze content addition patterns using time-series visualizations.",
    "Examine the distribution of movie durations.",
    "Investigate relationships between numerical variables.",
    "Communicate findings through clear visualizations and written explanations."
]:
    add_bullet(item)

# ------------------------------------------------------------
# 7. Dataset description
# ------------------------------------------------------------

add_heading("3. Dataset Description")

add_paragraph(
    f"The dataset contains {total_titles:,} records and {total_columns} "
    "columns. Important fields include show_id, type, title, director, "
    "cast, country, date_added, release_year, rating, duration, listed_in, "
    "and description."
)

add_paragraph(
    "The type column identifies movies and TV shows. The country column "
    "describes associated countries, date_added records the date a title "
    "was added according to the dataset, release_year indicates its "
    "release year, duration contains movie runtimes or TV show seasons, "
    "and listed_in contains genre and category labels."
)

add_paragraph(
    "The dataset is a historical snapshot rather than a complete record "
    "of Netflix's current catalog. Conclusions therefore describe the "
    "available dataset and should not automatically be generalized to "
    "Netflix's present-day content."
)

# ------------------------------------------------------------
# 8. Data exploration
# ------------------------------------------------------------

add_heading("4. Data Exploration and Preparation")

add_paragraph(
    f"Initial inspection identified {total_titles:,} records, "
    f"{total_columns} columns, and "
    f"{df.duplicated().sum():,} exact duplicate rows."
)

add_paragraph(
    f"The dataset contains {movie_count:,} movies and "
    f"{tv_count:,} TV shows. Movies account for "
    f"{movie_count / total_titles * 100:.2f}% of the titles, "
    f"while TV shows account for "
    f"{tv_count / total_titles * 100:.2f}%."
)

add_paragraph(
    f"The original dataset contains {total_missing:,} missing cells. "
    "Missing values are concentrated in fields such as director, cast, "
    "and country. Rather than removing every incomplete row, the analysis "
    "handles missing values according to the requirements of each chart."
)

add_paragraph(
    f"Movie-duration analysis uses {len(valid_durations):,} valid "
    f"duration values. The mean duration is {mean_duration:.2f} minutes "
    f"and the median is {median_duration:.2f} minutes."
)

add_heading("Data Preparation Methods", level=2)

for item in [
    "Loaded the CSV file using Pandas.",
    "Inspected data types, dimensions, missing values, and duplicates.",
    "Converted date_added into a datetime format for time-based analysis.",
    "Extracted numerical movie durations from the duration field.",
    "Split country and genre strings to count category appearances.",
    "Filtered records with missing values where required for individual analyses."
]:
    add_bullet(item)

# ------------------------------------------------------------
# 9. Visualization strategy
# ------------------------------------------------------------

add_heading("5. Visualization Strategy")

add_paragraph(
    "Different chart types were selected according to the analytical "
    "questions. Bar charts support categorical comparisons, line charts "
    "show changes over time, scatter plots reveal relationships between "
    "numerical variables, heatmaps summarize correlations, donut charts "
    "communicate proportions, and histograms display numerical distributions."
)

add_paragraph(
    "The charts use descriptive titles, axis labels, appropriate color "
    "schemes, legends where necessary, and saved high-resolution images. "
    "Color and formatting are intended to improve readability rather "
    "than exaggerate differences."
)

# ------------------------------------------------------------
# 10. Visualization findings
# ------------------------------------------------------------

add_heading("6. Visualizations and Findings")

add_chart(
    "01_top_10_countries.png",
    "6.1 Top 10 Countries by Netflix Content",
    "A horizontal bar chart compares the number of catalog titles "
    "associated with each of the ten leading countries. Country labels "
    "were split so that titles associated with multiple countries could "
    "contribute to each relevant country count.",
    f"The United States leads with {country_counts.get('United States', 0):,} "
    f"country-associated title appearances, followed by India with "
    f"{country_counts.get('India', 0):,}. These counts are not mutually "
    "exclusive because some titles are associated with multiple countries."
)

add_chart(
    "02_content_additions_by_year.png",
    "6.2 Annual Content Addition Trends",
    "A line chart displays the number of titles grouped by the year "
    "recorded in date_added. The visualization helps identify broad "
    "changes in catalog additions within the available snapshot.",
    f"The highest annual count in the dataset occurs in {peak_year}, "
    f"with {peak_additions:,} titles. This represents the peak in the "
    "available records, not necessarily Netflix's complete historical "
    "addition activity."
)

add_chart(
    "03_monthly_content_additions.png",
    "6.3 Monthly Content Addition Trends",
    "Monthly counts provide a more detailed view of the recorded "
    "addition dates. A three-month rolling average helps smooth "
    "short-term fluctuations and highlight broader patterns.",
    "Monthly peaks and fluctuations should be interpreted as patterns "
    "in the available records. Missing dates and the historical scope "
    "of the dataset may affect the observed pattern."
)

add_chart(
    "04_movie_duration_vs_release_year.png",
    "6.4 Movie Duration Versus Release Year",
    "Each point represents a movie, with release year on the horizontal "
    "axis and duration in minutes on the vertical axis. A smoothed "
    "yearly median line helps summarize changes across release years.",
    "The visualization allows comparisons of movie runtimes across "
    "release periods. The number of titles varies by year, so patterns "
    "in years with few observations should be treated cautiously."
)

add_chart(
    "09_movie_correlation_heatmap.png",
    "6.5 Movie Correlation Heatmap",
    "The corrected heatmap uses movie records only, ensuring that the "
    "duration variable represents minutes rather than TV show seasons. "
    "Pearson correlation coefficients summarize linear relationships "
    "between release year, year added, and movie duration.",
    f"Release year and duration have a correlation of "
    f"{correlation_matrix.loc['release_year', 'duration_minutes']:.3f}. "
    f"Year added and duration have a correlation of "
    f"{correlation_matrix.loc['year_added', 'duration_minutes']:.3f}. "
    "These are weak associations and do not establish causation."
)

add_chart(
    "06_content_distribution_donut.png",
    "6.6 Movies Versus TV Shows",
    "A donut chart compares the proportions of movies and TV shows "
    "in the available dataset. The total number of titles is shown "
    "in the center.",
    f"Movies account for {movie_count / total_titles * 100:.2f}% "
    f"of the dataset, while TV shows account for "
    f"{tv_count / total_titles * 100:.2f}%. These figures describe "
    "title counts, not audience viewing time or popularity."
)

add_chart(
    "07_top_10_genres.png",
    "6.7 Most Frequent Genres and Categories",
    "A horizontal bar chart ranks the ten most frequent category labels "
    "after splitting the listed_in field. A title can belong to "
    "multiple categories, so category counts can overlap.",
    f"The most frequent category is {genre_counts.index[0]}, with "
    f"{genre_counts.iloc[0]:,} appearances, followed by "
    f"{genre_counts.index[1]} with {genre_counts.iloc[1]:,}. "
    "These results describe category representation rather than "
    "exclusive groups of titles."
)

add_chart(
    "08_movie_duration_distribution.png",
    "6.8 Distribution of Movie Durations",
    "A histogram groups movie runtimes into intervals. Reference lines "
    "show the mean and median, allowing the central tendency and spread "
    "of movie durations to be assessed.",
    f"The median runtime is {median_duration:.0f} minutes and the "
    f"average is {mean_duration:.1f} minutes. The largest duration "
    "category in the prepared analysis is 91–120 minutes, containing "
    "2,996 movies."
)

# ------------------------------------------------------------
# 11. Key findings
# ------------------------------------------------------------

add_heading("7. Key Findings")

findings = [
    f"The dataset contains {total_titles:,} titles, of which "
    f"{movie_count:,} are movies and {tv_count:,} are TV shows.",

    f"The United States has {country_counts.get('United States', 0):,} "
    "country-associated title appearances, while India has "
    f"{country_counts.get('India', 0):,}.",

    f"The peak year in the recorded annual addition counts is {peak_year}, "
    f"with {peak_additions:,} titles.",

    f"The median movie duration is {median_duration:.0f} minutes, "
    f"and the average is {mean_duration:.1f} minutes.",

    f"The most frequent category is {genre_counts.index[0]}, with "
    f"{genre_counts.iloc[0]:,} category appearances.",

    "The corrected movie-only correlation analysis indicates a weak "
    "negative association between release year and movie duration."
]

for finding in findings:
    add_bullet(finding)

# ------------------------------------------------------------
# 12. Recommendations
# ------------------------------------------------------------

add_heading("8. Recommendations")

for item in [
    "Use consistent chart titles, labels, and formatting to improve readability.",
    "Investigate country and genre patterns further by comparing movies and TV shows separately.",
    "Use additional historical datasets before drawing conclusions about long-term changes in Netflix's catalog.",
    "Investigate unusual movie durations before deciding whether they represent errors or legitimate titles.",
    "Avoid interpreting correlation as evidence of causation.",
    "Consider interactive Plotly visualizations when users need to explore individual categories or observations."
]:
    add_bullet(item)

# ------------------------------------------------------------
# 13. Limitations
# ------------------------------------------------------------

add_heading("9. Limitations")

for item in [
    "The dataset represents a historical snapshot and does not necessarily reflect the current Netflix catalog.",
    "Missing values may affect country comparisons and analyses involving dates.",
    "Titles can be associated with multiple countries and genres, so some counts overlap.",
    "Correlation coefficients describe linear relationships and do not establish causation.",
    "The number of movies differs across release years, which can affect comparisons.",
    "The dataset does not directly measure viewership, audience satisfaction, or commercial success."
]:
    add_bullet(item)

# ------------------------------------------------------------
# 14. Conclusion
# ------------------------------------------------------------

add_heading("10. Conclusion")

add_paragraph(
    "This project demonstrates how Python-based data visualization "
    "can transform a raw entertainment dataset into an organized "
    "analytical report. Bar charts, line charts, scatter plots, "
    "heatmaps, donut charts, and histograms were used to examine "
    "catalog composition, content addition dates, genres, and movie "
    "durations. The results highlight the importance of choosing "
    "appropriate visualizations, preparing data carefully, and "
    "interpreting results in context. The project also demonstrates "
    "that clear visual storytelling must be supported by accurate "
    "calculations, transparent limitations, and cautious conclusions."
)

# ------------------------------------------------------------
# 15. Code documentation
# ------------------------------------------------------------

add_heading("11. Python Implementation")

add_paragraph(
    "The analysis was implemented using separate Python scripts. "
    "Pandas was used to load, filter, group, and summarize the dataset. "
    "Matplotlib was used to customize and save charts, while Seaborn "
    "was used to create the correlation heatmap. The following "
    "representative code illustrates the approach used for correlation "
    "analysis."
)

add_code_snippet(
    "correlation_matrix = numeric_df.corr(method='pearson')"
)

add_paragraph(
    "The full implementation is maintained in the project's notebooks "
    "folder. The chart images are stored in outputs/charts, and this "
    "report is saved in outputs/reports."
)

# ------------------------------------------------------------
# 16. References
# ------------------------------------------------------------

add_heading("12. References")

add_paragraph(
    "Dataset: Netflix Movies and TV Shows dataset, distributed through "
    "Kaggle by Shivam Bansal."
)

add_paragraph(
    "Dataset URL: https://www.kaggle.com/datasets/shivamb/netflix-shows"
)

add_paragraph(
    "Python documentation: https://docs.python.org/3/"
)

add_paragraph(
    "Pandas documentation: https://pandas.pydata.org/docs/"
)

add_paragraph(
    "Matplotlib documentation: https://matplotlib.org/stable/"
)

add_paragraph(
    "Seaborn documentation: https://seaborn.pydata.org/"
)

# ------------------------------------------------------------
# 17. Save report
# ------------------------------------------------------------

doc.save(REPORT_PATH)

print("\n" + "=" * 60)
print("WEEK 3 REPORT GENERATED")
print("=" * 60)

print(f"Dataset records: {total_titles:,}")
print(f"Charts folder: {CHART_DIR}")
print(f"Report saved to: {REPORT_PATH}")

print("\nPlease open the Word document and review its layout.")